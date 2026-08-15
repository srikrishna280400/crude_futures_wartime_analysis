"""_xlsx_helper.py — Helper to read .csv files that are actually XLSX.

This repo stores XLSX files with a .csv extension. openpyxl rejects the .csv
suffix explicitly, and pandas.read_csv cannot handle XLSX. This helper copies
the file to a tempfile with a .xlsx extension and reads it.

Stable location: scripts/_xlsx_helper.py (importable from other engines).
"""

from __future__ import annotations

import os
import shutil
import tempfile
from pathlib import Path

import pandas as pd


def read_csv_as_xlsx(path: str | Path, nrows: int | None = None) -> pd.DataFrame:
    """Read a file that is XLSX-format but saved as .csv."""
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    try:
        shutil.copy(str(path), tmp.name)
        return pd.read_excel(tmp.name, nrows=nrows)
    finally:
        os.unlink(tmp.name)


def get_headers(path: str | Path) -> list:
    from openpyxl import load_workbook  # local import to avoid hard dep at module load

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    try:
        shutil.copy(str(path), tmp.name)
        wb = load_workbook(tmp.name, read_only=True, data_only=True)
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(max_row=1))]
        wb.close()
    finally:
        os.unlink(tmp.name)
    return headers


def get_row_count(path: str | Path) -> int:
    from openpyxl import load_workbook

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    try:
        shutil.copy(str(path), tmp.name)
        wb = load_workbook(tmp.name, read_only=True, data_only=True)
        ws = wb.active
        n = ws.max_row - 1  # minus header
        wb.close()
    finally:
        os.unlink(tmp.name)
    return n
